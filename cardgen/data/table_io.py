"""Експорт таблиці даних у Excel та імпорт з Excel / CSV."""
from __future__ import annotations

import csv
import datetime as dt
import io
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .dataset import DataSet
from .fields import KEY_COLUMN, TYPE_TITLES, Schema, normalize_column
from .photos import Downloader, PhotoError, download, fetch_photo, store_image
from .xlsx_images import embedded_images

SHEET_TITLE = "Товари"
HEADER_FILL = PatternFill("solid", fgColor="1E5BFF")
REQUIRED_FILL = PatternFill("solid", fgColor="0B1F44")
EXCEL_ERRORS = {"#VALUE!", "#REF!", "#N/A", "#NAME?", "#DIV/0!", "#NUM!", "#NULL!"}


class TableImportError(Exception):
    pass


@dataclass
class ImportResult:
    dataset: DataSet
    warnings: list[str] = field(default_factory=list)
    photos: int = 0


# --- Експорт -----------------------------------------------------------------

def export_xlsx(schema: Schema, dataset: DataSet, path: Path) -> Path:
    """Таблиця для заповнення: стовпці шаблону + наявні дані."""
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_TITLE
    columns = schema.columns + [c for c in dataset.columns
                                if schema.get(c) is None and normalize_column(c) != normalize_column(KEY_COLUMN)]
    for col, name in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        spec = schema.get(name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = REQUIRED_FILL if spec and spec.required else HEADER_FILL
        cell.alignment = Alignment(vertical="center")
        cell.comment = Comment(_header_note(schema, name), "Генератор карток")
        ws.column_dimensions[get_column_letter(col)].width = 28 if spec and spec.is_image else 20
        if spec and spec.field_type == "list" and spec.options:
            formula = '"' + ",".join(o.replace('"', "'") for o in spec.options) + '"'
            if len(formula) <= 255:  # обмеження Excel для списку в правилі
                validation = DataValidation(type="list", formula1=formula, allow_blank=True)
                ws.add_data_validation(validation)
                validation.add(f"{get_column_letter(col)}2:{get_column_letter(col)}2000")
    for r, row in enumerate(dataset.rows, start=2):
        for col, name in enumerate(columns, start=1):
            ws.cell(row=r, column=col, value=dataset.value(row, name))
    ws.freeze_panes = "B2"
    ws.row_dimensions[1].height = 24

    help_sheet = wb.create_sheet("Як заповнювати")
    lines = [
        "Один рядок = один товар (одна картка).",
        f"«{KEY_COLUMN}» — обов'язковий і унікальний: за ним програма впізнає товар.",
        "Темно-сині заголовки — обов'язкові стовпці.",
        "Наведіть мишу на заголовок, щоб побачити підказку до стовпця.",
        "Фото можна вказати трьома способами:",
        "  1) вставити картинку прямо в клітинку (або поверх неї);",
        "  2) написати ім'я файлу (наприклад cable-18.jpg) — під час імпорту програма",
        "     попросить папку з фото;",
        "  3) вставити посилання https://… — програма завантажить фото.",
    ]
    for i, line in enumerate(lines, start=1):
        help_sheet.cell(row=i, column=1, value=line)
    help_sheet.column_dimensions["A"].width = 100
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _header_note(schema: Schema, name: str) -> str:
    spec = schema.get(name)
    if spec is None:
        return "Цей стовпець не використовується в шаблоні."
    parts = [f"Тип: {TYPE_TITLES.get(spec.field_type, spec.field_type)}"]
    if spec.required:
        parts.append("Обов'язкове")
    if spec.options:
        parts.append("Варіанти: " + ", ".join(spec.options))
    if spec.is_image:
        parts.append("Картинка в клітинці, ім'я файлу або посилання")
    parts.append("Використано: " + ", ".join(spec.used_in[:5]))
    return "\n".join(parts)


# --- Імпорт ------------------------------------------------------------------

def cell_text(value) -> str:
    """Значення клітинки -> текст, як його бачить людина в українському Excel."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "так" if value else "ні"
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return repr(value).replace(".", ",")
    if isinstance(value, dt.datetime):
        return value.strftime("%d.%m.%Y") if value.time() == dt.time() else value.strftime("%d.%m.%Y %H:%M")
    if isinstance(value, dt.date):
        return value.strftime("%d.%m.%Y")
    return str(value).strip()


def _read_xlsx(path: Path) -> tuple[list[list[str]], int]:
    """Рядки таблиці як текст і номер рядка заголовка (з 1)."""
    try:
        wb = load_workbook(path, data_only=True)
    except Exception as e:  # noqa: BLE001 — openpyxl кидає різні винятки
        raise TableImportError(f"Не вдалося відкрити {path.name}: {e}") from e
    ws = wb.worksheets[0]
    rows = [[cell_text(v) for v in row] for row in ws.iter_rows(values_only=True)]
    header = next((i for i, row in enumerate(rows) if any(row)), None)
    if header is None:
        raise TableImportError("Таблиця порожня")
    return rows[header:], header + 1


def _read_csv(path: Path) -> list[list[str]]:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise TableImportError("Невідоме кодування CSV-файлу")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t")
    except csv.Error:
        dialect = csv.excel
    rows = [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), dialect)]
    rows = [r for r in rows if any(r)]
    if not rows:
        raise TableImportError("Таблиця порожня")
    return rows


def import_table(path: Path, schema: Schema, images: Path, *, photos_folder: Path | None = None,
                 downloader: Downloader | None = None, previous: DataSet | None = None) -> ImportResult:
    """Читає таблицю. previous — поточні дані: якщо в новій таблиці фото товару не вказано,
    а раніше воно було (наприклад, додане в програмі), фото лишається."""
    result = _import(path, schema, images, photos_folder, downloader or download)
    if previous is not None:
        for row in result.dataset.rows:
            old = previous.find(result.dataset.key_of(row)) if result.dataset.key_of(row) else None
            if old is None:
                continue
            for spec in schema.fields:
                old_value = previous.value(old, spec.column)
                if spec.is_image and not row.get(spec.column) and old_value and (images / old_value).is_file():
                    row[spec.column] = old_value
    return result


def _import(path: Path, schema: Schema, images: Path, photos_folder: Path | None,
            downloader: Downloader) -> ImportResult:
    path = Path(path)
    is_csv = path.suffix.lower() == ".csv"
    if is_csv:
        rows, header_row = _read_csv(path), 1
        pictures = {}
    else:
        rows, header_row = _read_xlsx(path)
        pictures = embedded_images(path)

    headers = rows[0]
    warnings: list[str] = []
    columns: list[str] = []  # назва стовпця для кожного індексу файлу ("" — пропустити)
    for name in headers:
        spec = schema.get(name) if name else None
        if spec is not None:
            columns.append(spec.column)
        elif name:
            columns.append(name)
            warnings.append(f"Стовпець «{name}» не використовується в шаблоні — дані збережено, "
                            "але на картках їх не буде")
        else:
            columns.append("")
    if KEY_COLUMN not in columns:
        raise TableImportError(f"У таблиці немає стовпця «{KEY_COLUMN}»")
    present = {normalize_column(c) for c in columns if c}
    for spec in schema.fields:
        if normalize_column(spec.column) not in present:
            warnings.append(f"У таблиці немає стовпця «{spec.column}»")

    dataset = DataSet(columns=list(schema.columns), source=path.name)
    dataset.ensure_columns([c for c in columns if c])
    search_dirs = [d for d in (photos_folder, path.parent, images) if d is not None]
    photos = 0
    for offset, values in enumerate(rows[1:], start=1):
        excel_row = header_row + offset
        if not any(values):
            continue
        row = {c: "" for c in dataset.columns}
        for index, column in enumerate(columns):
            if column:
                row[column] = values[index] if index < len(values) else ""
        article = row[KEY_COLUMN].strip() or f"рядок-{excel_row}"
        for index, column in enumerate(columns):
            spec = schema.get(column) if column else None
            if spec is None or not spec.is_image:
                continue
            value = row[column]
            data = pictures.get((excel_row, index + 1))
            try:
                if data is None and value and value not in EXCEL_ERRORS:
                    data = fetch_photo(value, search_dirs, downloader)
                if data is not None:
                    row[column] = store_image(data, images, article, column)
                    photos += 1
                elif value in EXCEL_ERRORS:
                    row[column] = ""
                    warnings.append(f"Рядок {excel_row}: картинку в клітинці «{column}» не вдалося прочитати")
            except PhotoError as e:
                warnings.append(f"Рядок {excel_row}, «{column}»: {e}")
        dataset.rows.append(row)
    if not is_csv and pictures and photos == 0:
        warnings.append("У файлі є картинки, але жодна не стоїть у стовпці фото")
    return ImportResult(dataset, warnings, photos)
