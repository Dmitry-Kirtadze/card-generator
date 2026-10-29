"""Етап 6, інтерфейс: налаштування «Заповнюване», екран даних, сценарій «10 кабелів»."""
import io
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XlImage
from PIL import Image
from PySide6.QtWidgets import QCheckBox, QComboBox, QLineEdit

from cardgen.data.dataset import load_dataset
from cardgen.data.fields import KEY_COLUMN, build_schema
from cardgen.model import FieldBinding, PhotoElement, TableElement, TextElement
from cardgen.ui.editor.page import EditorPage
from cardgen.ui.main_window import MainWindow
from test_app import FakeDialogs


def png(color, path: Path | None = None, fmt="PNG") -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (60, 40), color).save(buffer, fmt)
    if path is not None:
        path.write_bytes(buffer.getvalue())
    return buffer.getvalue()


def visible(panel, cls, name):
    found = [w for w in panel.findChildren(cls, name) if w.isVisibleTo(panel)]
    assert len(found) == 1, (name, len(found))
    return found[0]


@pytest.fixture
def fake(monkeypatch):
    return FakeDialogs(monkeypatch)


@pytest.fixture
def cable(workspace):
    """Шаблон кабелю: назва, переріз (у таблиці), колір (список), фото."""
    c = workspace.create_category("Кабель AWG")
    t = workspace.create_template(c.id, "Обкладинка")
    table = TableElement.create(2, 2, name="Характеристики")
    table.cells[0][0].text, table.cells[1][0].text = "Переріз", "Колір"
    table.cells[0][1].binding = FieldBinding(enabled=True, column="Переріз", suffix=" AWG", required=True)
    table.cells[1][1].binding = FieldBinding(enabled=True, column="Колір", field_type="list",
                                             options=["червоний", "чорний"])
    t.elements = [
        TextElement(name="Заголовок", binding=FieldBinding(enabled=True, column="Назва", required=True)),
        PhotoElement(name="Фото"),
        table,
    ]
    workspace.save_template(c.id, t)
    return c, t


# --- Панель властивостей ----------------------------------------------------------

def test_binding_settings_in_panel(qtbot, workspace, cable):
    c, t = cable
    editor = EditorPage(workspace)
    qtbot.addWidget(editor)
    editor.show()
    editor.open_template(c.id, t.id)
    text = editor.add_block("text")
    panel = editor.properties
    assert visible(panel, QCheckBox, "bind_enabled").isChecked() is False
    assert not [w for w in panel.findChildren(QComboBox, "bind_column") if w.isVisibleTo(panel)]

    visible(panel, QCheckBox, "bind_enabled").setChecked(True)
    column = visible(panel, QComboBox, "bind_column")
    assert [column.itemText(i) for i in range(column.count())] == ["Назва", "Фото", "Переріз", "Колір"]
    column.setEditText("Довжина")
    kind = visible(panel, QComboBox, "bind_type")
    kind.setCurrentIndex(kind.findData("list"))
    options = visible(panel, QLineEdit, "bind_options")
    options.setText("1 м, 2 м")
    options.textEdited.emit("1 м, 2 м")
    visible(panel, QCheckBox, "bind_required").setChecked(True)
    b = text.binding
    assert (b.enabled, b.column, b.field_type, b.options, b.required) == \
        (True, "Довжина", "list", ["1 м", "2 м"], True)
    assert "Довжина" in build_schema(editor.template).columns
    editor.undo()
    assert not ops_find(editor, text.id).binding.required


def ops_find(editor, element_id):
    from cardgen.model import ops
    return ops.find(editor.template, element_id)


def test_photo_binding_can_be_turned_off(qtbot, workspace, cable):
    c, t = cable
    editor = EditorPage(workspace)
    qtbot.addWidget(editor)
    editor.show()
    editor.open_template(c.id, t.id)
    photo = editor.template.elements[1]
    editor.select_ids([photo.id])
    box = visible(editor.properties, QCheckBox, "bind_enabled")
    assert box.isChecked()
    box.setChecked(False)
    assert not photo.binding.enabled
    assert "Фото" not in build_schema(editor.template).columns


# --- Екран даних ------------------------------------------------------------------

@pytest.fixture
def window(qtbot, workspace):
    w = MainWindow(workspace)
    qtbot.addWidget(w)
    w.show()
    return w


def test_data_page_edit_add_delete(window, fake, workspace, cable, qtbot):
    c, t = cable
    window.open_data(c.id, t.id)
    page = window.data_page
    assert window.stack.currentWidget() is page
    assert [page.table.horizontalHeaderItem(i).text() for i in range(page.table.columnCount())] == \
        ["Артикул *", "Назва *", "Фото", "Переріз *", "Колір"]

    page.add_row()
    page.table.item(0, 0).setText("AWG-18-R")
    page.table.item(0, 1).setText("Кабель силіконовий 18 AWG")
    page.table.item(0, 3).setText("18")
    assert "помилок немає" in page.summary.text()
    page.table.item(0, 4).setText("синій")  # не зі списку — попередження
    assert "попереджень: 1" in page.summary.text()

    qtbot.waitUntil(lambda: load_dataset(workspace.tables_dir(c.id, t.id)).rows != [], timeout=3000)
    saved = load_dataset(workspace.tables_dir(c.id, t.id))
    assert saved.rows[0]["Назва"] == "Кабель силіконовий 18 AWG"

    page.table.selectRow(0)
    page.delete_rows()
    assert page.dataset.rows == []


def test_data_page_photo_by_drop(window, fake, workspace, cable, tmp_path):
    c, t = cable
    window.open_data(c.id, t.id)
    page = window.data_page
    page.add_row()
    page.table.item(0, 0).setText("A1")
    photo = tmp_path / "кабель.jpg"
    png((200, 0, 0), photo, "JPEG")
    page.table.fileDropped.emit(0, 2, str(photo))
    assert page.dataset.rows[0]["Фото"] == "A1_Фото.jpg"
    assert not page.table.item(0, 2).icon().isNull()
    page.table.fileDropped.emit(0, 1, str(photo))  # не стовпець фото
    assert fake.errors and "стовпець фото" in fake.errors[-1]


def test_validation_shown(window, fake, workspace, cable):
    c, t = cable
    window.open_data(c.id, t.id)
    page = window.data_page
    page.add_row()
    page.add_row()
    page.table.item(0, 0).setText("A1")
    page.table.item(1, 0).setText("A1")
    assert "помилок" in page.summary.text() and "немає" not in page.summary.text()
    messages = [page.issue_list.item(i).text() for i in range(page.issue_list.count())]
    assert any("повторюється" in m for m in messages)
    assert any("Обов'язкове" in m for m in messages)
    page._goto_issue(page.issue_list.item(0))
    assert page.table.currentRow() in (0, 1)


def test_export_from_data_page(window, fake, workspace, cable, tmp_path):
    c, t = cable
    window.open_data(c.id, t.id)
    fake.save_path = tmp_path / "заповнити"
    window.data_page.export_dialog()
    wb = load_workbook(tmp_path / "заповнити.xlsx")
    assert [cell.value for cell in wb["Товари"][1]] == [KEY_COLUMN, "Назва", "Фото", "Переріз", "Колір"]


# --- Приймання етапу: 10 кабелів --------------------------------------------------------

COLORS = ["червоний", "чорний"]


def fill_ten_cables(path: Path, photos_dir: Path) -> None:
    """Таблиця, як її заповнить людина: частина фото вставлена в Excel,
    частина — іменами файлів, одне — посиланням."""
    wb = Workbook()
    ws = wb.active
    ws.append([KEY_COLUMN, "Назва", "Фото", "Переріз", "Колір"])
    for i in range(10):
        awg = 10 + i * 2
        photo_cell = None
        if i < 4:
            photo_cell = f"awg{awg}.png"  # ім'я файлу
            png((20 * i, 100, 150), photos_dir / photo_cell)
        elif i == 9:
            photo_cell = f"https://shop.test/awg{awg}.png"
        ws.append([f"AWG-{awg}", f"Кабель силіконовий {awg} AWG", photo_cell, awg, COLORS[i % 2]])
    for i in range(4, 9):  # картинки, вставлені в Excel
        source = photos_dir / f"embedded{i}.png"
        png((200, 20 * i, 0), source)
        ws.add_image(XlImage(str(source)), f"C{i + 2}")
    wb.save(path)


def test_acceptance_ten_cables(window, fake, workspace, cable, tmp_path, monkeypatch):
    c, t = cable
    photos = tmp_path / "фото кабелів"
    photos.mkdir()
    xlsx = tmp_path / "кабелі.xlsx"
    fill_ten_cables(xlsx, photos)
    from cardgen.data import table_io
    monkeypatch.setattr(table_io, "download", lambda url: png((0, 0, 0)))

    # імпорт з екрана категорії, фото-файли лежать не поруч із таблицею — програма спитає папку
    window.show_category(c.id)
    window.category_page.grid.setCurrentRow(0)
    fake.open_path = xlsx
    fake.dir_path = photos
    (tmp_path / "awg10.png").unlink(missing_ok=True)
    moved = tmp_path / "інша"
    moved.mkdir()
    xlsx = xlsx.rename(moved / "кабелі.xlsx")
    fake.open_path = xlsx
    window.category_page.add_tables()

    page = window.data_page
    assert window.stack.currentWidget() is page
    assert any("папку" in text for text in fake.confirms)  # запитали папку з фото
    assert len(page.dataset.rows) == 10
    assert "помилок немає" in page.summary.text(), [i.message for i in page.issues]
    images = workspace.tables_dir(c.id, t.id) / "images"
    assert len(list(images.iterdir())) == 10  # усі фото різні — 10 файлів
    assert page.dataset.find("AWG-28")["Переріз"] == "28"
    assert "Імпортовано товарів: 10, фото: 10" in fake.infos[-1]

    # таблиця потрапила в історію завантажень, на екрані категорії видно підсумок
    page.leave()
    assert window.stack.currentWidget() is window.category_page
    assert "Товарів: 10 · без помилок" in window.category_page.data_summary.text()
    assert window.category_page.tables_list.count() == 1


def test_same_photo_for_many_products_stored_once(window, fake, workspace, cable, tmp_path):
    c, t = cable
    photo = tmp_path / "загальне.png"
    png((9, 9, 9), photo)
    wb = Workbook()
    wb.active.append([KEY_COLUMN, "Назва", "Фото", "Переріз"])
    for i in range(5):
        wb.active.append([f"A{i}", "Кабель", "загальне.png", 18])
    path = tmp_path / "кабелі.xlsx"
    wb.save(path)
    window.open_data(c.id, t.id)
    page = window.data_page
    assert page.import_file(path)
    images = workspace.tables_dir(c.id, t.id) / "images"
    assert [p.name for p in images.iterdir()] == ["A0_Фото.png"]
    assert {r["Фото"] for r in page.dataset.rows} == {"A0_Фото.png"}
    # видалили товари — фото лишається, поки ним користується хоч один
    page.table.selectRow(0)
    page.delete_rows()
    page.save_now()
    assert [p.name for p in images.iterdir()] == ["A0_Фото.png"]
    page.dataset.rows.clear()
    page.save_now()
    assert list(images.iterdir()) == []


def test_reimport_keeps_photos_added_in_app(window, fake, workspace, cable, tmp_path):
    c, t = cable
    window.open_data(c.id, t.id)
    page = window.data_page
    page.add_row()
    page.table.item(0, 0).setText("A1")
    photo = tmp_path / "p.png"
    png((1, 2, 3), photo)
    page.table.fileDropped.emit(0, 2, str(photo))
    wb = Workbook()
    wb.active.append([KEY_COLUMN, "Назва", "Переріз"])
    wb.active.append(["A1", "Нова назва", 20])
    path = tmp_path / "оновлення.xlsx"
    wb.save(path)
    assert page.import_file(path)
    row = page.dataset.find("A1")
    assert row["Назва"] == "Нова назва" and row["Фото"] == "A1_Фото.png"
